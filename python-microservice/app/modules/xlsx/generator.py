import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.schemas import GenerateReportRequest
from app.modules.common.logger import get_logger
from pathlib import Path

logger = get_logger(__name__)

class XLSXGenerator:
    def __init__(self):
        self.temp_dir = Path("temp")
        self.temp_dir.mkdir(exist_ok=True)

    async def generate_report(self, request: GenerateReportRequest) -> Path:
        """
        Generate an Excel report from the given request.
        """
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = request.reportType

            # Styles
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
            center_alignment = Alignment(horizontal="center", vertical="center")
            thin_border = Border(
                left=Side(style='thin'), 
                right=Side(style='thin'), 
                top=Side(style='thin'), 
                bottom=Side(style='thin')
            )

            # Title
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(request.columns))
            ws.cell(row=1, column=1).value = request.title
            ws.cell(row=1, column=1).font = Font(size=14, bold=True)
            ws.cell(row=1, column=1).alignment = center_alignment

            # Company Name (optional)
            if request.companyName:
                ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(request.columns))
                ws.cell(row=2, column=1).value = request.companyName
                ws.cell(row=2, column=1).font = Font(size=12)
                ws.cell(row=2, column=1).alignment = center_alignment
                header_row = 4
            else:
                header_row = 3

            # Headers
            for col_idx, column_name in enumerate(request.columns, 1):
                cell = ws.cell(row=header_row, column=col_idx)
                cell.value = column_name
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.alignment = center_alignment
                cell.border = thin_border

            # Data
            for row_idx, row_data in enumerate(request.data, header_row + 1):
                for col_idx, cell_value in enumerate(row_data, 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.value = cell_value
                    cell.border = thin_border
                    
                    # Auto-width adjustment (basic)
                    column_letter = ws.cell(row=row_idx, column=col_idx).column_letter
                    current_width = ws.column_dimensions[column_letter].width or 10
                    new_width = max(current_width, len(str(cell_value)) + 2)
                    ws.column_dimensions[column_letter].width = min(new_width, 50)

            # Generate filename
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"report_{request.reportType.lower()}_{timestamp}.xlsx"
            filepath = self.temp_dir / filename

            wb.save(filepath)
            logger.info(f"Report generated: {filepath}")
            return filepath

        except Exception as e:
            logger.error(f"Error generating XLSX report: {str(e)}", exc_info=True)
            raise
